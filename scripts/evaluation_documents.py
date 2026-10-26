"""Render the public, invented evaluation cases; no private CVs are accessed."""

import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
CASES_PATH = ROOT / "data/synthetic/extraction-eval/cases.json"


def cases() -> list[dict]:
    return json.loads(CASES_PATH.read_text(encoding="utf-8"))["cases"]


def render(case: dict, extension: str) -> bytes:
    stream = io.BytesIO()
    if extension == "docx":
        document = Document()
        document.styles["Normal"].font.size = Pt(11)
        if case["layout"] == "columns":
            for line in case["header"]:
                document.add_paragraph(line)
            table = document.add_table(rows=1, cols=2)
            for column, key in enumerate(["left", "right"]):
                table.cell(0, column).text = "\n".join(case[key])
        else:
            for line in case["lines"]:
                document.add_paragraph(line)
        document.save(stream)
    else:
        pdf = canvas.Canvas(stream, pagesize=A4, invariant=1)
        pdf.setTitle(f"{case['id']} - SYNTHETIC evaluation")
        pdf.setAuthor("RADIX ProofPath synthetic fixture generator")
        pdf.setFont("Helvetica", 10)
        if case["layout"] == "columns":
            for index, line in enumerate(case["header"]):
                pdf.drawString(36, 800 - index * 18, line)
            for x, key in [(36, "left"), (315, "right")]:
                for index, line in enumerate(case[key]):
                    pdf.drawString(x, 710 - index * 24, line)
        else:
            for index, line in enumerate(case["lines"]):
                pdf.drawString(36, 800 - index * 22, line)
        pdf.save()
    return stream.getvalue()


def write_pack(output: Path):
    output.mkdir(parents=True, exist_ok=True)
    for case in cases():
        for extension in case["formats"]:
            (output / f"{case['id']}.{extension}").write_bytes(render(case, extension))


if __name__ == "__main__":
    write_pack(ROOT / "data/synthetic/extraction-eval/documents")
