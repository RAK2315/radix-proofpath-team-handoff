"""Author labelled reconstructed fixtures; labels below precede extraction, never model predictions."""

import json
import os
import sys
from pathlib import Path
from textwrap import wrap

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
LABEL = "SYNTHETIC RECONSTRUCTION - NOT OFFICIAL ORGANIZER DATA"
SPECS = [
    {
        "id": "google-software-engineer",
        "kind": "jd",
        "company": "Google",
        "role": "Software Engineer",
        "required": [
            "Java",
            "Data Structures",
            "Algorithms",
            "Object-Oriented Design",
            "Git",
            "Testing",
            "System Design",
            "Communication",
        ],
        "preferred": ["Python", "Linux"],
    },
    {
        "id": "google-data-scientist",
        "kind": "jd",
        "company": "Google",
        "role": "Data Scientist",
        "required": ["Python", "SQL", "Machine Learning", "Statistics", "Pandas", "NumPy"],
        "preferred": ["PyTorch", "AWS"],
    },
    {
        "id": "microsoft-software-engineer",
        "kind": "jd",
        "company": "Microsoft",
        "role": "Software Engineer",
        "required": ["C#", "Object-Oriented Design", "REST APIs", "Testing", "SQL", "Git"],
        "preferred": ["Docker", "System Design"],
    },
    {
        "id": "microsoft-data-analyst",
        "kind": "jd",
        "company": "Microsoft",
        "role": "Data Analyst",
        "required": ["SQL", "Excel", "Power BI", "Statistics", "Communication"],
        "preferred": ["Python", "Pandas"],
    },
    {
        "id": "oracle-associate-software-engineer",
        "kind": "jd",
        "company": "Oracle Financial Services",
        "role": "Associate Software Engineer",
        "required": ["Java", "SQL", "Git", "Object-Oriented Design"],
        "preferred": ["Linux", "Testing"],
    },
    {
        "id": "oracle-application-support-analyst",
        "kind": "jd",
        "company": "Oracle Financial Services",
        "role": "Application Support Analyst",
        "required": ["SQL", "Linux", "Networking", "Communication"],
        "preferred": ["Shell Scripting", "TCP/IP"],
    },
    {
        "id": "ananya-rao",
        "kind": "resume",
        "name": "Ananya Rao",
        "orientation": "Systems-oriented generalist",
        "skills": [
            "Java",
            "Python",
            "C++",
            "Data Structures",
            "Algorithms",
            "Object-Oriented Design",
            "Git",
            "Linux",
            "Testing",
            "REST APIs",
            "System Design",
            "SQL",
            "Docker",
            "Communication",
        ],
        "project": "Built Java REST APIs with SQL storage and unit testing. Used Git and Docker on Linux.",
        "ratings": [7, 7, 6, 6, 5, 2, 5, 5, 7, 6, 5, 6],
    },
    {
        "id": "karthik-subramaniam",
        "kind": "resume",
        "name": "Karthik Subramaniam",
        "orientation": "Data-oriented analyst",
        "skills": [
            "Python",
            "SQL",
            "Pandas",
            "NumPy",
            "Machine Learning",
            "Statistics",
            "Power BI",
            "Excel",
            "PyTorch",
            "Git",
            "Communication",
        ],
        "project": "Built Python and SQL reports using Pandas, NumPy, Power BI and Excel. Evaluated machine learning experiments with PyTorch.",
        "ratings": [6, 4, 3, 7, 6, 7, 4, 8, 4, 3, 3, 4],
    },
    {
        "id": "priya-menon",
        "kind": "resume",
        "name": "Priya Menon",
        "orientation": "Entry-level support-oriented candidate",
        "skills": ["SQL", "Linux", "Networking", "Communication", "Shell Scripting", "TCP/IP", "Excel"],
        "project": "Prepared SQL troubleshooting notes and Linux checks. Used shell scripting and TCP/IP exercises to practice support tasks.",
        "ratings": [3, 2, 2, 5, 7, 1, 3, 5, 3, 2, 5, 5],
    },
    {
        "id": "rohan-verma",
        "kind": "resume",
        "name": "Rohan Verma",
        "orientation": "Broad all-rounder",
        "skills": [
            "Java",
            "Python",
            "C++",
            "C#",
            "Data Structures",
            "Algorithms",
            "Object-Oriented Design",
            "Git",
            "Linux",
            "Testing",
            "REST APIs",
            "System Design",
            "SQL",
            "Docker",
            "Communication",
            "Pandas",
            "NumPy",
            "Machine Learning",
            "Statistics",
            "Power BI",
            "Excel",
            "PyTorch",
            "AWS",
            "Networking",
            "Shell Scripting",
            "TCP/IP",
        ],
        "project": "Built Java and C# REST APIs with SQL, Git, unit testing and Docker. Created Python analytics using Pandas, NumPy, machine learning and AWS.",
        "ratings": [8, 8, 7, 8, 7, 7, 7, 8, 8, 7, 7, 7],
    },
]


def lines(spec):
    if spec["kind"] == "jd":
        return [
            LABEL,
            "Invented teaching text. No claim that this is a real vacancy or company policy.",
            f"Company: {spec['company']} (synthetic reconstruction)",
            f"Role: {spec['role']}",
            "About this synthetic role",
            "Practice the workflow described below in a fictional training environment.",
            "Required qualifications",
            *[f"- {name}" for name in spec["required"]],
            "Preferred qualifications",
            *[f"- {name}" for name in spec["preferred"]],
            "Responsibilities",
            "Contribute to a fictional team project and explain the original source of each requirement.",
        ]
    return [
        LABEL,
        f"Name: {spec['name']} (fictional reconstruction)",
        f"{spec['id']}@example.test",
        "Summary",
        spec["orientation"] + ". This identity and all experience are invented.",
        "Technical skills",
        ", ".join(spec["skills"]),
        "Projects",
        spec["project"],
        "Education",
        "Fictional Training Institute - computing studies (invented).",
        "Internships",
        "Synthetic practice placement - learning project only, not real employment.",
    ]


def create():
    from docx import Document
    from docx.shared import Pt
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    destination = ROOT / "data/synthetic/reconstructed"
    destination.mkdir(parents=True, exist_ok=True)
    # The labels are authored from the specification before any parser runs.
    labels = {
        "provenance": LABEL,
        "ground_truth": "Manually specified skills and importance from the authoring specification; not model-derived.",
        "cases": SPECS,
    }
    (destination / "expected-labels.json").write_text(json.dumps(labels, indent=2) + "\n", encoding="utf-8")
    for spec in SPECS:
        text = lines(spec)
        document = Document()
        document.styles["Normal"].font.size = Pt(11)
        for i, line in enumerate(text):
            if i == 0:
                document.add_heading(line, level=2)
            else:
                document.add_paragraph(line)
        document.save(destination / (spec["id"] + ".docx"))
        pdf = canvas.Canvas(str(destination / (spec["id"] + ".pdf")), pagesize=A4, invariant=1)
        pdf.setTitle(spec["id"] + " - " + LABEL)
        pdf.setAuthor("RADIX synthetic teaching fixture")
        y = 800
        for i, line in enumerate(text):
            pdf.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 9.1 if i == 0 else 10.5)
            for part in wrap(line, 85):
                if y < 45:
                    pdf.showPage()
                    y = 800
                    pdf.setFont("Helvetica", 10.5)
                pdf.drawString(36, y, part)
                y -= 15
            y -= 5
        pdf.save()
    (destination / "README.md").write_text(
        """# Reconstructed teaching pack

All ten sources are **SYNTHETIC RECONSTRUCTION - NOT OFFICIAL ORGANIZER DATA**.
Six JD roles and four fictional candidate names mirror the brief's descriptions, but these are invented
documents, not the missing original files, actual vacancies, company hiring policies or real people's CVs.
There are ten PDFs and ten DOCX files. `expected-labels.json` was specified before parser evaluation.
`computed/` contains actual rules-extraction results; these are predictions, not ground truth.
`matrix-results.json` records 24 computed candidate/JD scores and 12 synthetic company comparisons.
Candidate ratings in the specification are explicitly invented self-ratings, never inferred from CV text.
Use real previously unseen documents only with permission; do not commit personal input or model payloads.
""",
        encoding="utf-8",
    )
    print("Created 10 clearly labelled PDF/DOCX pairs with manually authored labels")


def evaluate():
    os.environ.update(RADIX_LLM_MODE="rules", RADIX_LLM_API_KEY="", GROQ_API_KEY="")
    from jd_analytics.backend.service import extract as jd_extract
    from resume_parsing.backend.service import extract as resume_extract
    from shared.contracts.models import CATEGORIES, Profile, Rating
    from shared.integrations.documents import read_document
    from shared.integrations.extraction import validate_evidence
    from skill_matching.backend.matcher import evaluate as match
    from talent_check.backend.benchmark_repository import load
    from talent_check.backend.scoring import evaluate as readiness

    destination = ROOT / "data/synthetic/reconstructed"
    computed = destination / "computed"
    computed.mkdir(exist_ok=True)
    jobs = {}
    people = {}
    observations = []
    for spec in SPECS:
        expected = set(spec.get("skills", spec.get("required", []) + spec.get("preferred", [])))
        for extension in ("pdf", "docx"):
            path = destination / (spec["id"] + "." + extension)
            sections = read_document(path.name, path.read_bytes())
            value = (jd_extract if spec["kind"] == "jd" else resume_extract)(sections, path.name, False)
            validate_evidence(value, sections)
            actual = {s.skill_name for s in value.skills}
            observations.append(
                {
                    "id": spec["id"],
                    "format": extension,
                    "missing": sorted(expected - actual),
                    "extra": sorted(actual - expected),
                    "quotes_valid": True,
                }
            )
            (computed / (spec["id"] + "-" + extension + ".json")).write_text(
                value.model_dump_json(indent=2) + "\n", encoding="utf-8"
            )
            if extension == "docx":
                value.review_status = "reviewed"
                if spec["kind"] == "jd":
                    jobs[spec["id"]] = value
                else:
                    people[spec["id"]] = Profile(
                        name=spec["name"],
                        skills=value.skills,
                        ratings=[
                            Rating(category_code=c, level=n)
                            for c, n in zip(CATEGORIES, spec["ratings"], strict=True)
                        ],
                    )
    matches = {
        name: {role: match(person, job)["match_score"] for role, job in jobs.items()}
        for name, person in people.items()
    }
    bundle = load("synthetic")
    company = {
        name: {c["id"]: readiness(person, c, bundle)["readiness_score"] for c in bundle["companies"]}
        for name, person in people.items()
    }
    data_jobs = ["google-data-scientist", "microsoft-data-analyst"]
    swe_jobs = [
        "google-software-engineer",
        "microsoft-software-engineer",
        "oracle-associate-software-engineer",
    ]
    data = matches["karthik-subramaniam"]
    checks = {
        "distinct_candidate_score_vectors": len({tuple(v.values()) for v in matches.values()}) == 4,
        "data_candidate_data_average_exceeds_swe_average": sum(data[j] for j in data_jobs) / 2
        > sum(data[j] for j in swe_jobs) / 3,
        "support_candidate_support_exceeds_google_swe": matches["priya-menon"][
            "oracle-application-support-analyst"
        ]
        > matches["priya-menon"]["google-software-engineer"],
    }
    report = {
        "provenance": LABEL,
        "mode": "Rules only; no live Groq or official data",
        "observations": observations,
        "matching": matches,
        "synthetic_readiness": company,
        "direction_checks": checks,
    }
    (destination / "matrix-results.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "matrix": matches,
                "direction_checks": checks,
                "observed_extraction_differences": [r for r in observations if r["missing"] or r["extra"]],
            },
            indent=2,
        )
    )
    assert all(checks.values()), "Report unexpected direction; do not change scorer to force rank"


if __name__ == "__main__":
    create()
    evaluate()
